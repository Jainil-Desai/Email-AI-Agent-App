import time
import random
from google.api_core.exceptions import TooManyRequests

class LLMResponder:
    def __init__(self, model, signature="Best,\nJainil Desai"):
        self.model = model
        self.signature = signature
        self.urgency_keywords = {
            'urgent': 5, 'asap': 5, 'immediately': 5, 'deadline': 4,
            'important': 4, 'critical': 4, 'emergency': 5, 'priority': 4,
            'time-sensitive': 4, 'action required': 3, 'please respond': 3
        }

    def _retry_generate(self, prompt, max_attempts=3):
        """Retry mechanism for rate-limited LLM requests"""
        for attempt in range(max_attempts):
            try:
                return self.model.generate_content(prompt).text.strip()
            except TooManyRequests:
                wait = 2 ** attempt + random.uniform(0, 1)
                time.sleep(wait)
        raise Exception("Rate limit exceeded")

    def analyze_email_priority(self, email_subject, email_body, sender):
        """Analyze email priority based on content and sender"""
        prompt = (
            "Analyze this email for urgency and importance. Consider:\n"
            "1. Subject line urgency indicators\n"
            "2. Content urgency indicators\n"
            "3. Sender's role/relationship\n"
            "4. Time-sensitive elements\n"
            "5. Action requirements\n\n"
            f"Subject: {email_subject}\n"
            f"From: {sender}\n"
            f"Content: {email_body[:1000]}\n\n"
            "Return a JSON with:\n"
            "- urgency_score (1-5)\n"
            "- importance_score (1-5)\n"
            "- reason (brief explanation)\n"
            "- suggested_response_time (immediate/within_hour/within_day/this_week)"
        )
        
        try:
            response = self._retry_generate(prompt)
            # Extract JSON from response
            import json
            import re
            json_match = re.search(r'\{.*\}', response, re.DOTALL)
            if json_match:
                return json.loads(json_match.group())
            return {
                'urgency_score': 1,
                'importance_score': 1,
                'reason': 'Unable to analyze priority',
                'suggested_response_time': 'this_week'
            }
        except Exception as e:
            return {
                'urgency_score': 1,
                'importance_score': 1,
                'reason': f'Error in priority analysis: {str(e)}',
                'suggested_response_time': 'this_week'
            }

    def detect_sentiment(self, text):
        """Enhanced sentiment analysis with detailed emotional tone detection"""
        prompt = (
            "Analyze the emotional tone of this message. Consider:\n"
            "1. Primary emotion\n"
            "2. Secondary emotions\n"
            "3. Emotional intensity (1-5)\n"
            "4. Key emotional triggers\n\n"
            "Emotions to consider: Joy, Anger, Disappointment, Anxiety, "
            "Frustration, Sadness, Neutral, Funny, Excitement, Concern, "
            "Gratitude, Urgency, Professional, Formal\n\n"
            f"Message: {text.strip()}\n\n"
            "Return a JSON with:\n"
            "- primary_emotion\n"
            "- secondary_emotions (array)\n"
            "- intensity (1-5)\n"
            "- triggers (array of key phrases)\n"
            "- emoji (most appropriate emoji)"
        )
        
        try:
            response = self._retry_generate(prompt)
            import json
            import re
            json_match = re.search(r'\{.*\}', response, re.DOTALL)
            if json_match:
                return json.loads(json_match.group())
            return {
                'primary_emotion': 'Neutral',
                'secondary_emotions': [],
                'intensity': 1,
                'triggers': [],
                'emoji': '😐'
            }
        except Exception as e:
            return {
                'primary_emotion': 'Neutral',
                'secondary_emotions': [],
                'intensity': 1,
                'triggers': [],
                'emoji': '😐'
            }

    def summarize_text(self, text, prefix="Summarize this text:", max_length=3000):
        """Enhanced text summarization with key points extraction"""
        # Truncate text if too long
        truncated_text = text.strip()[:max_length]
        
        prompt = (
            f"{prefix}\n\n"
            f"Text: {truncated_text}\n\n"
            "Provide a comprehensive summary including:\n"
            "1. Main points (bullet points)\n"
            "2. Key action items (if any)\n"
            "3. Important dates/deadlines (if any)\n"
            "4. Required responses/actions (if any)\n"
            "5. Context or background information\n\n"
            "Format the response with clear sections and bullet points."
        )
        
        return self._retry_generate(prompt) or "No summary available."

    def process_attachment(self, file_path, file_type):
        """Process different types of attachments with appropriate handling"""
        try:
            if file_type.lower() == 'pdf':
                return self._process_pdf(file_path)
            elif file_type.lower() in ['docx', 'doc']:
                return self._process_docx(file_path)
            elif file_type.lower() == 'txt':
                return self._process_txt(file_path)
            elif file_type.lower() == 'csv':
                return self._process_csv(file_path)
            elif file_type.lower() in ['xlsx', 'xls']:
                return self._process_excel(file_path)
            else:
                return f"[Unsupported file type: {file_type}]"
        except Exception as e:
            return f"[Error processing {file_type} file: {str(e)}]"

    def _process_pdf(self, file_path):
        """Process PDF files with text extraction and summarization"""
        try:
            import PyPDF2
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                text = ""
                for page in reader.pages:
                    text += page.extract_text() + "\n"
                return self.summarize_text(text, "Summarize this PDF document:")
        except Exception as e:
            return f"[Error processing PDF: {str(e)}]"

    def _process_docx(self, file_path):
        """Process DOCX files with text extraction and summarization"""
        try:
            import docx
            doc = docx.Document(file_path)
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            return self.summarize_text(text, "Summarize this Word document:")
        except Exception as e:
            return f"[Error processing DOCX: {str(e)}]"

    def _process_txt(self, file_path):
        """Process text files with summarization"""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                text = file.read()
                return self.summarize_text(text, "Summarize this text document:")
        except Exception as e:
            return f"[Error processing TXT: {str(e)}]"

    def _process_csv(self, file_path):
        """Process CSV files with data analysis and summarization"""
        try:
            import pandas as pd
            df = pd.read_csv(file_path)
            summary = f"CSV Summary:\nRows: {len(df)}\nColumns: {', '.join(df.columns)}\n"
            summary += "\nColumn Statistics:\n"
            for col in df.columns:
                if df[col].dtype in ['int64', 'float64']:
                    summary += f"{col}: min={df[col].min()}, max={df[col].max()}, mean={df[col].mean():.2f}\n"
            return summary
        except Exception as e:
            return f"[Error processing CSV: {str(e)}]"

    def _process_excel(self, file_path):
        """Process Excel files with data analysis and summarization"""
        try:
            import pandas as pd
            excel_file = pd.ExcelFile(file_path)
            summary = "Excel Summary:\n"
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                summary += f"\nSheet: {sheet_name}\n"
                summary += f"Rows: {len(df)}\nColumns: {', '.join(df.columns)}\n"
                summary += "Column Statistics:\n"
                for col in df.columns:
                    if df[col].dtype in ['int64', 'float64']:
                        summary += f"{col}: min={df[col].min()}, max={df[col].max()}, mean={df[col].mean():.2f}\n"
            return summary
        except Exception as e:
            return f"[Error processing Excel: {str(e)}]"

    def generate_reply_options(self, email_body, sender_name=None, num_options=3):
        """Generate personalized reply options with dynamic placeholders"""
        # Extract potential placeholders
        placeholders = self._extract_placeholders(email_body)
        
        prompt = (
            f"The following is an email that needs a reply:\n\n{email_body.strip()}\n\n"
            f"Generate {num_options} professional, friendly email replies.\n"
            f"Sender's name: {sender_name if sender_name else 'Not provided'}\n\n"
            "Requirements:\n"
            "1. Use appropriate greeting based on sender's name\n"
            "2. Include dynamic placeholders where personalization is needed:\n"
            "   - [NAME] for recipient's name\n"
            "   - [DATE] for specific dates\n"
            "   - [TIME] for specific times\n"
            "   - [COMPANY] for company names\n"
            "   - [DETAILS] for specific details to be filled\n"
            "3. Maintain professional tone\n"
            "4. Address all points from the original email\n"
            "5. End with signature\n\n"
            f"Signature: {self.signature}\n\n"
            "Format each option like:\n\n"
            "Subject: [Your subject here]\n"
            "Body: [Your reply here]\n\n"
            "Placeholders found in original email:\n" + 
            "\n".join([f"- {p}" for p in placeholders])
        )

        raw_output = self._retry_generate(prompt)
        blocks = raw_output.split("Subject:")[1:]

        options = []
        for block in blocks:
            lines = block.strip().splitlines()
            subject = lines[0].strip()
            remaining = "\n".join(lines[1:]).strip()
            if remaining.lower().startswith("body:"):
                remaining = remaining[5:].strip()

            if self.signature in remaining:
                body = remaining.split(self.signature)[0].strip() + f"\n{self.signature}"
            else:
                body = remaining.strip()

            options.append({
                "subject": subject,
                "body": body,
                "placeholders": self._extract_placeholders(body)
            })
        return options

    def _extract_placeholders(self, text):
        """Extract potential placeholders from text"""
        import re
        placeholder_patterns = {
            'name': r'\b(?:name|person|individual|recipient)\b',
            'date': r'\b(?:date|day|schedule|appointment|meeting)\b',
            'time': r'\b(?:time|hour|schedule|appointment|meeting)\b',
            'company': r'\b(?:company|organization|firm|business)\b',
            'details': r'\b(?:details|information|specifics|particulars)\b'
        }
        
        placeholders = set()
        for placeholder_type, pattern in placeholder_patterns.items():
            if re.search(pattern, text.lower()):
                placeholders.add(f"[{placeholder_type.upper()}]")
        
        return list(placeholders)

    def rank_emails_by_importance(self, email_summaries):
        """Enhanced email ranking based on multiple factors"""
        if not email_summaries:
            return []

        prompt = (
            "You are an email assistant. Rank these emails by urgency and importance.\n"
            "Consider:\n"
            "1. Sender's role/relationship\n"
            "2. Subject line urgency\n"
            "3. Content urgency\n"
            "4. Time-sensitive elements\n"
            "5. Action requirements\n\n"
            "Emails to rank:\n\n"
        )
        
        for i, email in enumerate(email_summaries, 1):
            prompt += (
                f"{i}. From: {email['from']}\n"
                f"Subject: {email['subject']}\n"
                f"Summary: {email['summary']}\n"
                f"Priority Analysis: {email.get('priority_analysis', 'Not analyzed')}\n"
                f"Sentiment: {email.get('sentiment', 'Not analyzed')}\n\n"
            )
        
        prompt += (
            "Return a JSON with:\n"
            "- ranked_indices (array of indices in order of priority, starting from 0)\n"
            "- reasons (array of reasons for each ranking)\n"
            "- suggested_actions (array of suggested actions for each email)\n\n"
            f"Note: Indices must be between 0 and {len(email_summaries) - 1}"
        )

        try:
            response = self._retry_generate(prompt)
            import json
            import re
            json_match = re.search(r'\{.*\}', response, re.DOTALL)
            if json_match:
                ranking_data = json.loads(json_match.group())
                indices = ranking_data.get('ranked_indices', [])
                
                # Validate indices
                valid_indices = []
                seen_indices = set()
                for idx in indices:
                    # Check if index is valid and not duplicate
                    if isinstance(idx, int) and 0 <= idx < len(email_summaries) and idx not in seen_indices:
                        valid_indices.append(idx)
                        seen_indices.add(idx)
                
                # If we have valid indices, return them
                if valid_indices:
                    return valid_indices
                
            # Fallback: return indices in original order
            return list(range(len(email_summaries)))
        except Exception as e:
            print(f"Error in email ranking: {str(e)}")
            return list(range(len(email_summaries))) 